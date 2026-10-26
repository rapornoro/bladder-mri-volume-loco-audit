"""Build submission Word files (European Radiology): Times New Roman 12, double spacing (main/title/ESM),
2.5 cm margins, centered page numbers, bordered single-spaced tables; checklists in landscape.
Usage: python s22_build_docx.py [submission_folder_name]   (default: submission_eurradiol)"""
import os
import sys
import subprocess
import copy
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

SUB = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "manuscript",
                                   sys.argv[1] if len(sys.argv) > 1 else "submission_eurradiol"))
REF_DIR = os.path.join(SUB, "_templates")
os.makedirs(REF_DIR, exist_ok=True)
FONT = "Times New Roman"


def page_number_footer(section):
    p = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)
    run.font.name = FONT; run.font.size = Pt(10)


def set_style_font(style, size, bold=None, spacing=None, after=None, color_black=True):
    style.font.name = FONT
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if bold is not None:
        style.font.bold = bold
    if color_black:
        style.font.color.rgb = RGBColor(0, 0, 0)
    if not hasattr(style, "paragraph_format"):
        return
    pf = style.paragraph_format
    if spacing is not None:
        pf.line_spacing = spacing
    if after is not None:
        pf.space_after = Pt(after); pf.space_before = Pt(0)


def make_reference(name, landscape, body_size, spacing):
    path = os.path.join(REF_DIR, name)
    with open(path, "wb") as fh:
        fh.write(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"],
                                capture_output=True, check=True).stdout)
    d = Document(path)
    for st in d.styles:
        if st.type == 1 and st.name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text",
                                        "Bibliography", "Abstract", "Date", "Author", "Caption", "Table Caption",
                                        "Image Caption", "Footnote Text"):
            set_style_font(st, body_size, spacing=spacing, after=6)
    for st in d.styles:  # inline code / code blocks (file names) in the journal font too
        if st.name in ("Verbatim Char", "Source Code"):
            set_style_font(st, body_size - 1, color_black=True)
    for st_name, size in (("Title", body_size + 2), ("Heading 1", body_size + 2), ("Heading 2", body_size),
                          ("Heading 3", body_size), ("Heading 4", body_size)):
        try:
            set_style_font(d.styles[st_name], size, bold=True, spacing=spacing, after=6)
            d.styles[st_name].font.italic = False
        except KeyError:
            pass
    for s in d.sections:
        if landscape:
            s.orientation = WD_ORIENT.LANDSCAPE
            s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        else:
            s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, m, Cm(2.5))
        page_number_footer(s)
    d.save(path)
    return path


def borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(b)


def postprocess(path, table_size):
    d = Document(path)
    for tbl in d.tables:
        borders(tbl)
        for ri, row in enumerate(tbl.rows):
            if ri == 0:  # repeat header row on each page
                trPr = row._tr.get_or_add_trPr()
                h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(table_size); r.font.name = FONT
                        if ri == 0:
                            r.font.bold = True
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    d.core_properties.comments = ""
    d.save(path)


def build(md, ref, table_size):
    out = os.path.join(SUB, md.replace(".md", ".docx"))
    subprocess.run(["pandoc", os.path.join(SUB, md), "-o", out, "--reference-doc", ref], check=True)
    postprocess(out, table_size)
    return out


if __name__ == "__main__":
    portrait2 = make_reference("ref_portrait_double.docx", False, 12, 2.0)
    portrait1 = make_reference("ref_portrait_single.docx", False, 12, 1.15)
    land = make_reference("ref_landscape.docx", True, 10, 1.0)
    files = sorted(f for f in os.listdir(SUB) if f.endswith(".md"))
    double = {"title_page.md"}
    def is_double(f):
        return f in double or f.startswith("manuscript") or f.startswith("supplement")
    skip = {"figure_plan.md", "submission_checklist.md"}
    for md in files:
        if md in skip:
            continue
        if md.startswith("checklist_") or md.startswith("table"):
            print(build(md, land, 9))
        elif is_double(md):
            print(build(md, portrait2, 9))
        else:
            print(build(md, portrait1, 10))
